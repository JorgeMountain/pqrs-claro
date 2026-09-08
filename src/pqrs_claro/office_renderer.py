"""Oficio canónico determinista y comprobación completa de su Word y PDF."""

from __future__ import annotations

import hashlib
import re
import unicodedata
from copy import deepcopy
from datetime import date, datetime, timezone
from pathlib import Path
from zipfile import ZipFile

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from pydantic import ValidationError

from .models import DomainError
from .office_models import (
    MAX_ATTACHMENT_BYTES,
    MAX_OFFICE_CHARS,
    MAX_OFFICE_PAGES,
    OfficeMetadata,
    xml_text,
)
from .office_templates import (
    image_pixels,
    template_profile,
    validate_letterhead,
    validate_template_file,
    verify_preserved_template,
)
from .office_word import convert_word, pdf_capabilities
from .review_models import ReviewParagraph
from .review_service import input_hash
from .source_library import digest
from .source_retrieval import source_in_scope

OFFICE_VERSION = 2
__all__ = ["OFFICE_VERSION", "build_document", "render_office", "pdf_capabilities"]
_MONTHS = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)


def _text(value, *, required=True) -> str:
    try:
        value = xml_text(value)
    except ValueError as exc:
        raise DomainError(str(exc)) from None
    if required and not value.strip():
        raise DomainError("El oficio contiene un texto obligatorio vacío.")
    return value


def _hash_valid(value: dict, key="snapshot_hash") -> bool:
    return (
        isinstance(value, dict)
        and isinstance(value.get(key), str)
        and value[key] == digest({name: item for name, item in value.items() if name != key})
    )


def _public_reference(source: dict) -> str:
    title = _text(source.get("title"))
    if re.search(r"(?:[A-Za-z]:[\\/]|file://)", title):
        raise DomainError(
            "Una fuente tiene una ruta local como título. Corrígela antes de exportar."
        )
    location = "sección" if source.get("location_kind") == "section" else "página"
    if any(
        type(source.get(key)) is not int or source[key] < 1
        for key in ("version", "page", "line_start", "line_end")
    ):
        raise DomainError("Una referencia aprobada no conserva su versión o ubicación.")
    if source["line_end"] < source["line_start"]:
        raise DomainError("Una referencia aprobada contiene un rango de líneas inválido.")
    _text(source.get("quote"))
    return (
        f"{title}. Versión {source['version']}, {location} {source['page']}, "
        f"líneas {source['line_start']}–{source['line_end']}."
    )


def _approved_snapshots(approved: dict) -> list[dict]:
    if not _hash_valid(approved) or approved.get("status") != "approved":
        raise DomainError("El expediente no tiene un conjunto aprobado íntegro.")
    snapshots = approved.get("requests")
    if not isinstance(snapshots, list) or not 1 <= len(snapshots) <= 200:
        raise DomainError("El oficio requiere entre 1 y 200 solicitudes aprobadas.")
    if approved.get("approved_count") != len(snapshots) or approved.get("total_count") != len(
        snapshots
    ):
        raise DomainError("Todas las solicitudes vigentes deben estar aprobadas.")
    first = snapshots[0]
    identification = first.get("identification") or {}
    request_order = [request.get("id") for request in identification.get("requests", [])]
    if len(set(request_order)) != len(snapshots) or request_order != [
        s.get("request_id") for s in snapshots
    ]:
        raise DomainError(
            "El orden de las solicitudes no corresponde a la identificación aprobada."
        )
    for snapshot in snapshots:
        if not _hash_valid(snapshot) or snapshot.get("case_id") != approved.get("case_id"):
            raise DomainError(
                "Una aprobación no corresponde a este expediente o perdió integridad."
            )
        if snapshot.get("identification") != identification or any(
            snapshot.get(key) != first.get(key) for key in ("case_version", "case_original_sha256")
        ):
            raise DomainError("Las aprobaciones pertenecen a versiones distintas del expediente.")
        review, approval = snapshot.get("review"), snapshot.get("approval")
        if not isinstance(review, dict) or not isinstance(approval, dict):
            raise DomainError("Una respuesta no conserva la revisión y su aprobación.")
        if (
            review.get("case_id") != approved.get("case_id")
            or review.get("request_id") != snapshot["request_id"]
        ):
            raise DomainError("Una respuesta aprobada pertenece a otra solicitud.")
        assessment = review.get("assessment") or {}
        fingerprint = input_hash(review)
        if (
            assessment.get("status") != "passed"
            or assessment.get("issues")
            or assessment.get("input_hash") != fingerprint
            or approval.get("input_hash") != fingerprint
            or snapshot.get("id") != approval.get("id")
            or review.get("stale")
            or any(check.get("passed") is not True for check in assessment.get("local_checks", []))
        ):
            raise DomainError(
                "La respuesta no conserva una evaluación aprobable de su contenido exacto."
            )
        _text(approval.get("approved_by"))
        _text(approval.get("approved_at"))
        if any(item.get("resolution") is None for item in review.get("pending", [])):
            raise DomainError("Una respuesta conserva pendientes sin resolver.")
        request = next(
            item for item in identification["requests"] if item["id"] == snapshot["request_id"]
        )
        recorded = review.get("request_snapshot") or {}
        if any(recorded.get(key) != value for key, value in request.items()):
            raise DomainError("La respuesta no corresponde a la solicitud identificada.")
        if recorded.get("case_id") != approved["case_id"]:
            raise DomainError("La identificación de la respuesta pertenece a otro expediente.")
        if any(review.get(key) != request.get(key) for key in ("label", "title", "description")):
            raise DomainError("El título o contenido de la solicitud cambió después de aprobar.")
    return snapshots


def build_document(
    approved, metadata_dict, attachments_list, prepared_by, *, letterhead=None
) -> dict:
    """Deriva únicamente contenido publicable de snapshots aprobados completos."""
    try:
        metadata = OfficeMetadata.model_validate(metadata_dict).model_dump(mode="json")
    except ValidationError:
        raise DomainError(
            "Completa los datos del oficio con valores válidos, sin campos adicionales."
        ) from None
    prepared_by = _text(prepared_by).strip()
    if len(prepared_by) > 240:
        raise DomainError("El nombre de quien prepara el oficio admite hasta 240 caracteres.")
    snapshots = _approved_snapshots(approved)
    references, reference_index, requests, cited_documents = [], {}, [], {}

    def reference(kind, key, text):
        if key not in reference_index:
            number = len(references) + 1
            reference_index[key] = number
            references.append({"number": number, "kind": kind, "text": _text(text)})
        return reference_index[key]

    for snapshot in snapshots:
        review = snapshot["review"]
        available = review.get("available_supports", [])
        facts = review.get("facts", [])
        if not isinstance(available, list) or not isinstance(facts, list):
            raise DomainError("La revisión no conserva un inventario válido de fuentes y aportes.")
        sources = {source["chunk_id"]: source for source in available}
        facts = {fact["id"]: fact for fact in facts}
        if len(sources) != len(available) or len(facts) != len(review.get("facts", [])):
            raise DomainError("Una respuesta conserva referencias con identificadores duplicados.")
        try:
            paragraphs = [
                ReviewParagraph.model_validate(paragraph).model_dump()
                for paragraph in review["paragraphs"]
            ]
        except (ValidationError, KeyError, TypeError):
            raise DomainError(
                "Una respuesta aprobada contiene párrafos o citas inválidos."
            ) from None
        if not 1 <= len(paragraphs) <= 30:
            raise DomainError("Cada respuesta aprobada necesita entre 1 y 30 párrafos.")
        rendered = []
        for paragraph in paragraphs:
            numbers = []
            if not paragraph["cited_chunk_ids"] and not paragraph["cited_fact_ids"]:
                raise DomainError("Un párrafo aprobado no conserva referencias.")
            for chunk_id in paragraph["cited_chunk_ids"]:
                source = sources.get(chunk_id)
                if not source or not source_in_scope(review["request_snapshot"], source):
                    raise DomainError(
                        "Un párrafo cita una fuente desconocida o fuera de su ámbito."
                    )
                if source.get("validity") == "superseded":
                    raise DomainError("Un párrafo cita una versión reemplazada de una fuente.")
                key = ("document", source["source_id"], source["version"], chunk_id)
                numbers.append(reference("document", key, _public_reference(source)))
                existing = cited_documents.get(source["source_id"])
                if existing and any(
                    existing.get(field) != source.get(field)
                    for field in ("sha256", "document_id", "version")
                ):
                    raise DomainError(
                        "Una fuente citada tiene versiones o huellas contradictorias."
                    )
                cited_documents[source["source_id"]] = source
            for fact_id in paragraph["cited_fact_ids"]:
                fact = facts.get(fact_id)
                if (
                    not fact
                    or fact.get("case_id") != review["case_id"]
                    or fact.get("request_id") != review["request_id"]
                ):
                    raise DomainError(
                        "Un párrafo cita una declaración retirada o de otra solicitud."
                    )
                if not _hash_valid(fact, "sha256"):
                    raise DomainError("Una declaración citada perdió su integridad.")
                _text(fact.get("text"))
                _text(fact.get("basis"))
                author = _text(fact.get("author"))
                try:
                    observed = date.fromisoformat(fact["observed_at"]).isoformat()
                except (ValueError, KeyError, TypeError):
                    raise DomainError(
                        "Una declaración citada no conserva su fecha de conocimiento."
                    ) from None
                numbers.append(
                    reference(
                        "fact",
                        ("fact", review["request_id"], fact_id),
                        f"Declaración de {author}. Fecha de conocimiento: {observed}.",
                    )
                )
            rendered.append({"text": _text(paragraph["text"]), "reference_numbers": numbers})
        requests.append(
            {
                "request_id": review["request_id"],
                "label": _text(review["label"]),
                "title": _text(review["title"]),
                "description": _text(review["description"]),
                "paragraphs": rendered,
            }
        )

    if not isinstance(attachments_list, list) or len(attachments_list) > 20:
        raise DomainError("El oficio admite como máximo 20 anexos.")
    attachments, seen, total_size = [], set(), 0
    for item in attachments_list:
        if not isinstance(item, dict) or set(item) != {
            "source_id",
            "title",
            "filename",
            "size",
            "sha256",
        }:
            raise DomainError("Un anexo no tiene los metadatos esperados.")
        source = cited_documents.get(item["source_id"])
        if not source or item["source_id"] in seen or item["sha256"] != source["sha256"]:
            raise DomainError("Solo pueden anexarse documentos citados, íntegros y sin duplicados.")
        if type(item["size"]) is not int or item["size"] <= 0:
            raise DomainError("Un anexo no tiene un tamaño válido.")
        for key in ("title", "filename"):
            _text(item[key])
        if item["filename"] != Path(item["filename"].replace("\\", "/")).name:
            raise DomainError("El nombre del anexo no debe contener una ruta.")
        if item["title"] != source["title"] or item["filename"] != source["filename"]:
            raise DomainError("El anexo no conserva el título y nombre de la fuente citada.")
        seen.add(item["source_id"])
        total_size += item["size"]
        attachments.append(deepcopy(item))
    if total_size > MAX_ATTACHMENT_BYTES:
        raise DomainError("Los anexos superan el máximo conjunto de 200 MiB.")
    document = {
        "metadata": metadata,
        "requests": requests,
        "references": references,
        "attachments": attachments,
        "prepared_by": prepared_by,
    }
    if letterhead is not None:
        document["letterhead"] = validate_letterhead(letterhead)
    _blocks(document)
    return document


def _blocks(document: dict) -> list[dict]:
    """Una sola secuencia define lo visible en el DOCX y lo exigido al PDF."""
    metadata = document["metadata"]
    issued = date.fromisoformat(metadata["issue_date"])
    blocks = []

    def add(kind, text, **extra):
        blocks.append({"kind": kind, "text": _text(text), **extra})

    if not document.get("letterhead"):
        add("institution", metadata["institution"])
        if metadata["department"]:
            add("department", metadata["department"])
        add(
            "date",
            f"{metadata['city']}, {issued.day} de {_MONTHS[issued.month - 1]} de {issued.year}",
        )
    if metadata["office_number"]:
        add("metadata", "Oficio " + metadata["office_number"])
    if document.get("letterhead"):
        add(
            "date",
            f"{metadata['city']}, {issued.day} de {_MONTHS[issued.month - 1]} de {issued.year}",
        )
    add("recipient", metadata["recipient"])
    if metadata["recipient_contact"]:
        add("metadata", metadata["recipient_contact"])
    add("subject", "Asunto: " + metadata["subject"])
    if metadata["petition_number"]:
        add("metadata", "Petición: " + metadata["petition_number"])
    add("intro", "En atención a su petición, se presentan las respuestas a cada solicitud.")
    known = {item["number"] for item in document["references"]}
    for request_index, request in enumerate(document["requests"]):
        add("heading", f"Solicitud {request['label']}. {request['title']}")
        add("request", request["description"])
        for paragraph_index, paragraph in enumerate(request["paragraphs"]):
            add(
                "response",
                paragraph["text"],
                bookmark=f"response_{request_index}_{paragraph_index}",
            )
            if (
                not paragraph["reference_numbers"]
                or not set(paragraph["reference_numbers"]) <= known
            ):
                raise DomainError("La vista previa contiene referencias inválidas.")
            add(
                "citation",
                "Referencias: " + ", ".join(f"[{n}]" for n in paragraph["reference_numbers"]) + ".",
            )
    add("closing", "Atentamente,")
    add("signer", metadata["signer_name"])
    add("signer_role", metadata["signer_role"])
    if document["references"]:
        add("heading", "Fuentes y declaraciones citadas")
        for reference in document["references"]:
            add("reference", f"[{reference['number']}] {reference['text']}")
    if document["attachments"]:
        add("heading", "Documentos incluidos en el paquete")
        for index, attachment in enumerate(document["attachments"], 1):
            add("attachment", f"{index}. {attachment['title']} ({attachment['filename']})")
    if sum(len(block["text"]) for block in blocks) > MAX_OFFICE_CHARS:
        raise DomainError(
            "El oficio supera el máximo de 600.000 caracteres. No se recortará su contenido."
        )
    return blocks


def _make_docx(blocks: list[dict], path: Path) -> None:
    document = Document()
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = "Respuesta a petición"
    properties.subject = ""
    properties.comments = ""
    fixed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    properties.created = properties.modified = fixed
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(0.9)
    section.header_distance = section.footer_distance = Inches(0.35)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Footer"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.08
        style.paragraph_format.widow_control = True
    document.styles["Title"].font.size = Pt(15)
    document.styles["Title"].font.bold = True
    document.styles["Heading 1"].font.size = Pt(12)
    document.styles["Heading 1"].font.bold = True
    document.styles["Heading 1"].paragraph_format.space_before = Pt(13)
    document.styles["Heading 1"].paragraph_format.keep_with_next = True
    # La plantilla de python-docx puede traer un borde azul y fuentes de tema
    # que Word prioriza sobre font.name. También neutralizamos el estilo enlazado.
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Footer"):
        style = document.styles[name]
        selected = [style]
        linked = style.element.find(qn("w:link"))
        if linked is not None:
            selected.extend(
                candidate
                for candidate in document.styles
                if candidate.style_id == linked.get(qn("w:val"))
            )
        for candidate in selected:
            candidate.font.name = "Times New Roman"
            candidate.font.color.rgb = RGBColor(0, 0, 0)
            candidate.font.size = style.font.size
            run_properties = candidate.element.get_or_add_rPr()
            fonts = run_properties.find(qn("w:rFonts"))
            for attribute in list(fonts.attrib):
                if "theme" in attribute.casefold():
                    del fonts.attrib[attribute]
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{attribute}"), "Times New Roman")
            for tag in ("w:spacing", "w:szCs"):
                for inherited in run_properties.findall(qn(tag)):
                    run_properties.remove(inherited)
            for border in candidate.element.findall(f"{qn('w:pPr')}/{qn('w:pBdr')}"):
                border.getparent().remove(border)
    for index, block in enumerate(blocks):
        kind = block["kind"]
        style = "Title" if kind == "institution" else "Heading 1" if kind == "heading" else "Normal"
        paragraph = document.add_paragraph(style=style)
        if "bookmark" in block:
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(index))
            start.set(qn("w:name"), block["bookmark"])
            paragraph._p.append(start)
        run = paragraph.add_run(block["text"])
        if "bookmark" in block:
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(index))
            paragraph._p.append(end)
        if kind in {"subject", "signer"}:
            run.bold = True
        if kind in {"citation", "reference", "attachment"}:
            run.font.size = Pt(10)
            paragraph.paragraph_format.space_after = Pt(8)
        if kind == "request":
            run.italic = True
        if kind in {"institution", "department", "subject", "closing", "signer"}:
            paragraph.paragraph_format.keep_with_next = True
        if kind == "closing":
            paragraph.paragraph_format.space_before = Pt(14)
        # Una sola palabra larga puede partirse, sin introducir guiones automáticos.
        suppress = OxmlElement("w:suppressAutoHyphens")
        paragraph._p.get_or_add_pPr().append(suppress)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Página ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(path)


def _make_template_docx(blocks: list[dict], path: Path, template_path: Path, profile: dict) -> None:
    """Reemplaza solo el cuerpo editable y los metadatos personales de la copia."""
    document = Document(template_path)
    body = document.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for index, block in enumerate(blocks):
        paragraph = document.add_paragraph()
        if profile["paragraph_properties"] is not None:
            paragraph._p.insert(0, deepcopy(profile["paragraph_properties"]))
        kind = block["kind"]
        # El cuerpo original emplea formato de carácter directo. Cambiar Normal
        # modificaría también el pie institucional que hereda ese estilo.
        run = paragraph.add_run(block["text"])
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        fonts = run._r.get_or_add_rPr().find(qn("w:rFonts"))
        for name in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{name}"), "Times New Roman")
        if kind in {"heading", "subject", "signer"}:
            run.bold = True
        if kind == "request":
            run.italic = True
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.widow_control = True
        if kind in {"heading", "subject", "closing", "signer"}:
            paragraph.paragraph_format.keep_with_next = True
        if kind not in {"response", "request"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if kind == "closing":
            paragraph.paragraph_format.space_before = Pt(12)
        if "bookmark" in block:
            start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
            start.set(qn("w:id"), str(index))
            start.set(qn("w:name"), block["bookmark"])
            end.set(qn("w:id"), str(index))
            paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
            paragraph._p.append(end)
        paragraph._p.get_or_add_pPr().append(OxmlElement("w:suppressAutoHyphens"))
    properties = document.core_properties
    properties.author = properties.last_modified_by = properties.subject = properties.comments = ""
    properties.title = "Respuesta a petición"
    properties.created = properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    replacements = {
        "word/document.xml": etree.tostring(
            document.element, xml_declaration=True, encoding="UTF-8", standalone=True
        ),
        "docProps/core.xml": etree.tostring(
            properties._element, xml_declaration=True, encoding="UTF-8", standalone=True
        ),
    }
    with ZipFile(template_path) as original, ZipFile(path, "w") as output:
        for item in original.infolist():
            output.writestr(item, replacements.get(item.filename, original.read(item.filename)))
    verify_preserved_template(template_path, path)


def _verify_docx(blocks: list[dict], path: Path) -> bool:
    docx = Document(path)
    if [paragraph.text for paragraph in docx.paragraphs] != [block["text"] for block in blocks]:
        raise DomainError("El Word generado no conserva exactamente todos los bloques aprobados.")
    expected = {block["bookmark"]: block["text"] for block in blocks if "bookmark" in block}
    found = {}
    for paragraph in docx.paragraphs:
        for marker in paragraph._p.findall(qn("w:bookmarkStart")):
            name = marker.get(qn("w:name"))
            if name in found:
                raise DomainError("El Word repite un identificador de párrafo aprobado.")
            found[name] = paragraph.text
    if found != expected:
        raise DomainError("No fue posible verificar cada párrafo aprobado en el Word.")
    return True


def _normalized(text: str) -> str:
    return " ".join(unicodedata.normalize("NFC", text).split())


def _pdf_body_text(page) -> str:
    lines, footers = [], 0
    for block in page.get_text("dict", sort=True)["blocks"]:
        for line in block.get("lines", []):
            text = "".join(span["text"] for span in line["spans"])
            x0, y0, x1, _ = line["bbox"]
            # Se excluye solo el número de página exacto en su posición fija.
            # Recortar una zona puede conservar medio glifo del pie o esconder
            # contenido adicional; toda otra línea sigue en la comparación.
            if (
                _normalized(text) == f"Página {page.number + 1}"
                and y0 >= page.rect.height - 52
                and abs((x0 + x1) / 2 - page.rect.width / 2) <= 20
            ):
                footers += 1
            else:
                lines.append(text)
    if footers > 1:
        raise DomainError("El PDF repite su pie de página; la conversión no es fiel.")
    return "\n".join(lines)


def _template_pdf_body_text(page, profile: dict) -> str:
    if (
        abs(page.rect.width - profile["page_width"]) > 1
        or abs(page.rect.height - profile["page_height"]) > 1
    ):
        raise DomainError("El PDF no conserva el tamaño de página del membrete.")
    header, footer = profile["regions"]
    parts = {"header": [], "footer": [], "body": []}
    for block in page.get_text("dict", sort=True)["blocks"]:
        for line in block.get("lines", []):
            text = "".join(span["text"] for span in line["spans"])
            y0 = line["bbox"][1]
            region = (
                "header"
                if y0 < header["boundary"]
                else "footer"
                if y0 >= footer["boundary"] - 2
                else "body"
            )
            parts[region].append(text)
    # La zona estática no se ignora: debe contener exactamente el texto aprobado
    # del membrete en cada página. Cualquier añadido, ausencia o cruce falla.
    for region in (header, footer):
        if _normalized("\n".join(parts[region["kind"]])) != _normalized(region["text"]):
            raise DomainError(
                "El PDF alteró u omitió el encabezado o pie institucional del membrete."
            )
    return "\n".join(parts["body"])


def _verify_template_images(page, profile: dict) -> None:
    images = page.get_image_info(xrefs=True)
    if len(images) != len(profile["images"]):
        raise DomainError("El PDF omitió o añadió una imagen del membrete institucional.")
    masks = {item[0]: item[1] for item in page.get_images(full=True)}
    used = set()
    for expected in profile["images"]:
        matched = False
        for index, actual in enumerate(images):
            if index in used or (actual["width"], actual["height"]) != (
                expected["width"],
                expected["height"],
            ):
                continue
            bbox = actual["bbox"]
            if expected["bbox"] is not None:
                if any(abs(a - b) > 1 for a, b in zip(bbox, expected["bbox"])):
                    continue
            elif expected["region"] == "header":
                if bbox[1] < 0 or bbox[3] > profile["regions"][0]["boundary"] + 1:
                    continue
            elif bbox[1] < profile["regions"][1]["boundary"] - 2 or bbox[3] > page.rect.height:
                continue
            if not actual["xref"]:
                continue
            pixmap = pymupdf.Pixmap(page.parent, actual["xref"])
            mask = masks.get(actual["xref"])
            if mask:
                if pixmap.alpha:
                    pixmap = pymupdf.Pixmap(pixmap, 0)
                pixmap = pymupdf.Pixmap(pixmap, pymupdf.Pixmap(page.parent, mask))
            pixels = image_pixels(pixmap)
            # Word mantiene resolución y píxeles del PNG; la composición de su
            # máscara introduce redondeos de como máximo 2 niveles sobre 255.
            if len(pixels) == len(expected["pixels"]) and all(
                abs(a - b) <= 2 for a, b in zip(pixels, expected["pixels"])
            ):
                used.add(index)
                matched = True
                break
        if not matched:
            raise DomainError("El PDF alteró el logo institucional o su posición en el membrete.")


def _pdf_text_matches(expected: str, actual: str) -> bool:
    expected = _normalized(expected)
    if _normalized(actual) == expected:
        return True
    lines = [_normalized(line) for line in actual.splitlines() if line.strip()]
    if not lines:
        return False
    pattern = re.escape(lines[0])
    for previous, current in zip(lines, lines[1:]):
        # Word puede cortar físicamente 2026-09-07 después del guion. Solo
        # permitimos ese espacio de extracción en límites reales de línea;
        # jamás eliminamos espacios internos ni modificamos un carácter.
        separator = " ?" if previous.endswith("-") and current[0].isalnum() else " "
        pattern += separator + re.escape(current)
    return re.fullmatch(pattern, expected) is not None


def _verify_pdf(blocks: list[dict], path: Path, *, profile=None) -> int:
    try:
        with pymupdf.open(path) as pdf:
            if not pdf.is_pdf or pdf.needs_pass or not 1 <= len(pdf) <= MAX_OFFICE_PAGES:
                raise DomainError("El PDF debe ser legible y tener entre 1 y 200 páginas.")
            pages = [
                _template_pdf_body_text(page, profile) if profile else _pdf_body_text(page)
                for page in pdf
            ]
            if profile:
                for page in pdf:
                    _verify_template_images(page, profile)
            actual = "\n".join(pages)
            expected = _normalized("\n".join(block["text"] for block in blocks))
            if not _pdf_text_matches(expected, actual):
                raise DomainError(
                    "El PDF no reproduce todos los bloques y referencias en su orden exacto. "
                    "No se habilitará la descarga de una conversión incompleta o alterada."
                )
            return len(pdf)
    except (pymupdf.FileDataError, RuntimeError, ValueError) as exc:
        if isinstance(exc, DomainError):
            raise
        raise DomainError("No fue posible comprobar el PDF producido por Word.") from None


def render_office(document, output_dir, *, converter=None, template_path=None) -> dict:
    """Escribe una vez; solo declara éxito después de cotejar Word y PDF completos."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    docx, pdf = output_dir / "oficio.docx", output_dir / "oficio.pdf"
    if docx.exists() or pdf.exists():
        raise DomainError("La generación no sobrescribe archivos de una versión anterior.")
    letterhead = document.get("letterhead")
    profile = None
    if bool(letterhead) != (template_path is not None):
        raise DomainError("La plantilla no corresponde al membrete revisado en la vista previa.")
    if letterhead:
        letterhead = validate_letterhead(letterhead)
        template_path = validate_template_file(template_path, letterhead["sha256"])
        profile = template_profile(template_path)
    blocks = _blocks(document)
    if profile:
        _make_template_docx(blocks, docx, template_path, profile)
    else:
        _make_docx(blocks, docx)
    _verify_docx(blocks, docx)
    if profile:
        validate_template_file(template_path, letterhead["sha256"])
        verify_preserved_template(template_path, docx)
    original_sha = hashlib.sha256(docx.read_bytes()).hexdigest()
    try:
        result = (converter or convert_word)(docx, pdf)
    except DomainError:
        raise
    except Exception:
        raise DomainError("No se pudo convertir el Word a PDF. El oficio no quedó listo.") from None
    if hashlib.sha256(docx.read_bytes()).hexdigest() != original_sha:
        raise DomainError("El conversor modificó el Word canónico; la generación fue descartada.")
    _verify_docx(blocks, docx)
    if profile:
        validate_template_file(template_path, letterhead["sha256"])
        verify_preserved_template(template_path, docx)
    if not pdf.is_file() or not pdf.stat().st_size:
        raise DomainError("El conversor no produjo un PDF válido; el oficio no quedó listo.")
    pages = _verify_pdf(blocks, pdf, profile=profile)
    return {
        "page_count": pages,
        "converter": result
        if isinstance(result, str)
        else ("Microsoft Word COM" if converter is None else "Conversor de prueba"),
        "verification": {
            "docx_exact": True,
            "pdf_text_match": True,
            "request_count": len(document["requests"]),
            "paragraph_count": sum(len(request["paragraphs"]) for request in document["requests"]),
        },
    }
