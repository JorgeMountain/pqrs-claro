"""Biblioteca local real con archivos ficticios; no hace llamadas a OpenAI."""

from __future__ import annotations

import hashlib
import json
from io import BytesIO

import pymupdf
import pytest
from docx import Document

from pqrs_claro.extraction import PDFExtractor
from pqrs_claro.models import DomainError
from pqrs_claro.source_library import SourceLibrary, digest


@pytest.fixture
def library(tmp_path):
    return SourceLibrary(tmp_path / "data", PDFExtractor())


def add_text(library, text="EJEMPLO FICTICIO: requisitos del trámite.", **overrides):
    metadata = {
        "title": "Reglas ficticias",
        "category": "institutional_norm",
        "origin": "Documento ficticio local",
        **overrides,
    }
    return library.add("reglas.txt", text.encode("utf-8"), **metadata)


def pdf_bytes(page_count):
    with pymupdf.open() as document:
        for number in range(1, page_count + 1):
            page = document.new_page()
            page.insert_text(
                (40, 60),
                f"FUENTE FICTICIA. Requisitos y procedimiento de prueba. Pagina {number}.",
            )
        return document.tobytes()


def test_collection_reading_counts_reflect_current_extraction_instead_of_import_totals(library):
    library.save_collection({"id":"archivo", "name":"Archivo ficticio",
                             "extraction_counts":{"pending":99}})
    source = add_text(library)
    source.update(collection_id="archivo", extraction_status="pending")
    with library.connect() as connection:
        connection.execute("UPDATE sources SET body=? WHERE id=?",
                           (json.dumps(source),source["id"]))
    assert library.collections()[0]["extraction_counts"] == {"pending":1}
    source["extraction_status"] = "ready"
    with library.connect() as connection:
        connection.execute("UPDATE sources SET body=? WHERE id=?",
                           (json.dumps(source),source["id"]))
    assert library.collections()[0]["extraction_counts"] == {"ready":1}


def test_txt_sections_keep_exact_text_and_all_chunk_references(library):
    lines = [
        f"  Apartado {number:02d}: información ficticia con acentos y espacios "
        "para comprobar citas."
        for number in range(1, 84)
    ]
    payload = ("\ufeff" + "\r\n".join(lines)).encode("utf-8")
    source = library.add(
        "reglas.TXT",
        payload,
        title="Reglas ficticias extensas",
        category="institutional_norm",
        origin="Fixture TXT UTF-8 con BOM",
    )
    result = library.get(source["id"])
    pages = result["pages"]
    assert source["location_kind"] == "section"
    assert source["page_count"] == 3
    assert [page["number"] for page in pages] == [1, 2, 3]
    assert "\n".join(page["text"] for page in pages) == "\n".join(lines)
    assert library.original(source["id"]).read_bytes() == payload
    assert source["sha256"] == hashlib.sha256(payload).hexdigest()

    sources, chunks = library.snapshot("case-a")
    assert [item["id"] for item in sources] == [source["id"]]
    assert len(chunks) > len(pages)
    covered = set()
    for chunk in chunks:
        page_lines = pages[chunk["page"] - 1]["text"].splitlines()
        assert chunk["quote"] == "\n".join(
            page_lines[chunk["line_start"] - 1 : chunk["line_end"]]
        )
        assert chunk["source_id"] == source["id"]
        assert chunk["location_kind"] == "section"
        covered.update(
            (chunk["page"], number)
            for number in range(chunk["line_start"], chunk["line_end"] + 1)
        )
    assert covered == {
        (page["number"], number)
        for page in pages
        for number in range(1, len(page["text"].splitlines()) + 1)
    }


def test_docx_preserves_paragraph_and_table_order(library):
    document = Document()
    document.add_paragraph("EJEMPLO FICTICIO: párrafo anterior a la tabla.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Dato"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "Estado"
    table.cell(1, 1).text = "En revisión"
    document.add_paragraph("Párrafo posterior: falta el archivo del acta.")
    output = BytesIO()
    document.save(output)
    source = library.add(
        "registro.docx",
        output.getvalue(),
        title="Registro ficticio",
        category="case_evidence",
        origin="Fixture DOCX",
        case_id="case-a",
    )
    result = library.get(source["id"], "case-a")
    assert source["location_kind"] == "section"
    assert result["pages"][0]["text"] == (
        "EJEMPLO FICTICIO: párrafo anterior a la tabla.\n"
        "Dato | Valor\n"
        "Estado | En revisión\n"
        "Párrafo posterior: falta el archivo del acta."
    )
    assert library.original(source["id"], "case-a").read_bytes() == output.getvalue()


@pytest.mark.parametrize("page_count", [41, 300])
def test_source_pdf_accepts_more_than_petition_limit(library, page_count):
    source = library.add(
        "norma.pdf",
        pdf_bytes(page_count),
        title="Norma ficticia extensa",
        category="national_norm",
        origin="Fixture PDF nativo",
    )
    pages = library.get(source["id"])["pages"]
    assert source["page_count"] == page_count
    assert source["location_kind"] == "page"
    assert [page["number"] for page in pages] == list(range(1, page_count + 1))
    assert f"Pagina {page_count}." in pages[-1]["text"]
    assert all(page["method"] == "native" and not page["needs_review"] for page in pages)


def test_source_pdf_rejects_above_300_before_reading_and_removes_upload(library, monkeypatch):
    def unexpected_read(page):
        pytest.fail("El límite debe comprobarse antes de leer páginas u ofrecer OCR.")

    monkeypatch.setattr(library.pdf_extractor, "_read_page", unexpected_read)
    with pytest.raises(DomainError, match="300 páginas"):
        library.add(
            "excesiva.pdf",
            pdf_bytes(301),
            title="Fuente demasiado extensa",
            category="national_norm",
            origin="Fixture PDF de 301 páginas",
        )
    assert library.list() == []
    assert list(library.root.iterdir()) == []


def test_new_version_keeps_previous_original_and_requires_its_own_review(library):
    old_text = "EJEMPLO FICTICIO. Requisito: constancia de inscripción."
    new_text = old_text + " Nuevo requisito: comprobante de actualización de datos."
    old = add_text(library, old_text)
    library.review(
        old["id"],
        reviewed_by="Responsable ficticio",
        review_note="Se contrastó el original ficticio de la primera versión.",
        validity="current",
    )
    updated = add_text(library, new_text, document_id=old["document_id"])
    assert updated["id"] != old["id"]
    assert updated["version"] == 2
    assert updated["document_id"] == old["document_id"]
    assert updated["review_status"] == "pending"
    assert updated["validity"] == "unknown"
    assert updated["reviewed_by"] is None
    assert not library.get(old["id"])["source"]["is_latest"]
    assert library.get(updated["id"])["source"]["is_latest"]
    assert library.original(old["id"]).read_text(encoding="utf-8") == old_text
    assert library.original(updated["id"]).read_text(encoding="utf-8") == new_text
    sources, chunks = library.snapshot("case-a")
    assert [source["id"] for source in sources] == [updated["id"]]
    assert {chunk["source_id"] for chunk in chunks} == {updated["id"]}
    repeated = add_text(library, new_text, document_id=old["document_id"])
    assert repeated["id"] == updated["id"]
    assert len(library.list()) == 2


def test_case_evidence_requires_scope(library):
    with pytest.raises(DomainError, match="asociarse a un expediente"):
        add_text(library, category="case_evidence")
    assert library.list() == []
    assert list(library.root.iterdir()) == []


def test_other_case_evidence_is_invisible_through_every_read_path(library):
    common = add_text(library)
    own = add_text(library, "Registro ficticio A.", category="case_evidence", case_id="case-a")
    other = add_text(library, "Registro ficticio B.", category="case_evidence", case_id="case-b")
    assert {source["id"] for source in library.list("case-a")} == {common["id"], own["id"]}
    assert {source["id"] for source in library.list()} == {common["id"]}
    for wrong_case in ("case-a", None):
        with pytest.raises(KeyError, match="este expediente"):
            library.get(other["id"], wrong_case)
        with pytest.raises(KeyError, match="este expediente"):
            library.original(other["id"], wrong_case)
    sources, chunks = library.snapshot("case-a")
    assert {source["id"] for source in sources} == {common["id"], own["id"]}
    assert {chunk["source_id"] for chunk in chunks} == {common["id"], own["id"]}


def test_review_changes_snapshot_and_superseded_source_leaves_candidates(library):
    source = add_text(library)
    before_sources, before_chunks = library.snapshot("case-a")
    reviewed = library.review(
        source["id"],
        reviewed_by="Responsable ficticio",
        review_note="Se comprobó la versión del ejemplo.",
        validity="current",
    )
    after_sources, after_chunks = library.snapshot("case-a")
    assert digest(before_sources) != digest(after_sources)
    assert digest(before_chunks) != digest(after_chunks)
    assert after_sources[0]["reviewed_by"] == "Responsable ficticio"
    assert reviewed["review_status"] == "reviewed"
    assert after_chunks[0]["quote"] == before_chunks[0]["quote"]
    assert after_chunks[0]["validity"] == "current"
    library.review(
        source["id"],
        reviewed_by="Otro responsable ficticio",
        review_note="Esta regla fue reemplazada dentro del ejemplo.",
        validity="superseded",
    )
    assert library.snapshot("case-a") == ([], [])
    assert library.original(source["id"]).is_file()


def test_previous_response_is_available_in_library_but_not_evidence_snapshot(library):
    norm = add_text(library)
    before = library.snapshot("case-a")
    previous = add_text(
        library,
        "EJEMPLO FICTICIO: se remitió copia del acta del caso distinto EJEMPLO-999.",
        category="previous_response",
        title="Oficio anterior ficticio para estilo",
    )
    assert {source["id"] for source in library.list()} == {norm["id"], previous["id"]}
    assert library.get(previous["id"])["pages"]
    assert library.snapshot("case-a") == before


def test_tampered_original_blocks_download_and_search_snapshot(library):
    source = add_text(library)
    original = library.original(source["id"])
    original.write_bytes(b"Contenido modificado fuera de la biblioteca.")
    with pytest.raises(DomainError, match="original cambió"):
        library.original(source["id"])
    with pytest.raises(DomainError, match="original cambió"):
        library.snapshot("case-a")
