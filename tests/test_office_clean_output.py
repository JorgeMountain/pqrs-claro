"""Salida institucional limpia con trazabilidad conservada fuera del texto impreso."""

import re
from copy import deepcopy

import pymupdf
import pytest
import test_office_templates as template_support
from docx import Document
from docx.oxml.ns import qn
from test_office_renderer import fake_pdf, metadata, reseal, source
from test_office_structure import (
    approved_with_original_refs,
    original_ref,
    paragraph,
    selected,
    structured,
)

from pqrs_claro.models import DomainError
from pqrs_claro.office_renderer import _blocks, build_document, render_office
from pqrs_claro.office_templates import load_letterhead

NATURAL_REFERENCE = (
    "El Oficio 00995-2026 distingue los períodos 2015-1 y 2019-2. "
    "La mención del artículo 25 del Acuerdo Superior 083 de 1996 "
    "se conserva con su alcance aprobado."
)


def clean_approval():
    context = paragraph("El antecedente se limita a los períodos consultados.", "context")
    value = structured(
        [context, paragraph(NATURAL_REFERENCE)],
        [
            context,
            paragraph("El segundo numeral conserva la distinción documental."),
            paragraph("La conclusión se limita al segundo numeral.", "conclusion"),
        ],
    )
    # Marcadores que solo pueden venir de la referencia interna, nunca de la prosa.
    for snapshot in value["requests"]:
        support = snapshot["review"]["available_supports"][0]
        support.update(title="LOCALIZADOR INTERNO DE PRUEBA", page=91, line_start=72, line_end=74)
    return reseal(value)


def assert_clean(text):
    assert not re.search(r"\[\d+\]", text)
    for internal in (
        "Referencias",
        "Fuentes consultadas",
        "Antecedente para los numerales",
        "Sobre el numeral",
        "Fragmento de la petición",
        "LOCALIZADOR INTERNO DE PRUEBA",
        "NOTA INTERNA NO PUBLICABLE",
        "Candidato privado no citado",
        "source-1:1:1-3",
    ):
        assert internal not in text


def test_hidden_reference_blocks_keep_every_approved_origin_and_citation():
    value = clean_approval()
    before = deepcopy(value)
    document = build_document(value, metadata(), [], "Preparador privado")
    assert value == before
    assert not {"citation", "reference", "scope"} & {
        block["kind"] for block in document["blocks"]
    }
    assert_clean("\n".join(block["text"] for block in document["blocks"]))
    assert len(document["references"]) == 1
    assert "LOCALIZADOR INTERNO DE PRUEBA" in document["references"][0]["text"]
    contexts = selected(document, "context")
    assert len(contexts) == 1
    assert contexts[0]["request_ids"] == ["P1", "P2"]
    assert contexts[0]["bookmarks"] == ["response_0_0", "response_1_0"]
    for i, request in enumerate(document["requests"]):
        approved_paragraphs = before["requests"][i]["review"]["paragraphs"]
        assert [p["text"] for p in request["paragraphs"]] == [
            p["text"] for p in approved_paragraphs
        ]
        for j, item in enumerate(request["paragraphs"]):
            assert item["reference_numbers"] == [1]
            assert item["bookmark"] == f"response_{i}_{j}"
            block = next(
                b for b in document["blocks"] if item["bookmark"] in b.get("bookmarks", [])
            )
            assert block["reference_numbers"] == [1]
            assert request["request_id"] in block["request_ids"]
            assert approved_paragraphs[j]["cited_chunk_ids"] == [source()["chunk_id"]]


@pytest.mark.parametrize("with_template", [False, True])
def test_word_and_pdf_omit_technical_apparatus_but_preserve_text_and_bookmarks(
    tmp_path, with_template
):
    value = clean_approval()
    before = deepcopy(value)
    options = {}
    render_options = {"converter": fake_pdf}
    if with_template:
        path, _, data_dir = template_support.template.__wrapped__(tmp_path)
        options["letterhead"] = load_letterhead(data_dir).public()
        render_options = {"template_path": path, "converter": template_support.template_pdf(path)}
    document = build_document(value, metadata(), [], "Preparador privado", **options)
    canonical_before = deepcopy(document)
    output = tmp_path / "export"
    result = render_office(document, output, **render_options)
    assert result["verification"]["docx_exact"] is True
    assert result["verification"]["pdf_text_match"] is True
    word = Document(output / "oficio.docx")
    body = "\n".join(p.text for p in word.paragraphs)
    with pymupdf.open(output / "oficio.pdf") as pdf:
        pdf_text = "\n".join(page.get_text() for page in pdf)
    assert_clean(body)
    assert_clean(pdf_text)
    assert NATURAL_REFERENCE in body
    assert " ".join(NATURAL_REFERENCE.split()) in " ".join(pdf_text.split())
    bookmarks = word.element.findall(f".//{qn('w:bookmarkStart')}")
    names = [item.get(qn("w:name")) for item in bookmarks]
    expected = [p["bookmark"] for request in document["requests"] for p in request["paragraphs"]]
    assert len(names) == len(expected) == len(set(names))
    assert set(names) == set(expected)
    for block in document["blocks"]:
        if block.get("bookmarks"):
            matching = next(p for p in word.paragraphs if p.text == block["text"])
            assert {
                marker.get(qn("w:name")) for marker in matching._p.findall(qn("w:bookmarkStart"))
            } == set(block["bookmarks"])
    assert value == before
    assert document == canonical_before


def test_discontinuous_request_quotes_print_omission_without_source_coordinates(tmp_path):
    refs = [
        original_ref(91, 72, 72, "Segundo fragmento literal."),
        original_ref(8, 39, 40, "Primer fragmento literal.\nContinúa la pregunta."),
    ]
    value = approved_with_original_refs(refs)
    before = deepcopy(value)
    document = build_document(value, metadata(), [], "Preparador privado")
    requests = selected(document, "request")
    assert [b["text"] for b in requests] == [
        "Primer fragmento literal.\nContinúa la pregunta.",
        "Segundo fragmento literal.",
    ]
    assert [b["original_location"] for b in requests] == [
        {"page": 8, "line_start": 39, "line_end": 40},
        {"page": 91, "line_start": 72, "line_end": 72},
    ]
    assert [b["text"] for b in selected(document, "omission")] == ["[…]"]
    assert all(b["request_ids"] == ["P1"] for b in requests)
    assert len(document["requests"]) == len(selected(document, "response")) == 1
    render_office(document, tmp_path, converter=fake_pdf)
    body = "\n".join(p.text for p in Document(tmp_path / "oficio.docx").paragraphs)
    with pymupdf.open(tmp_path / "oficio.pdf") as pdf:
        pdf_text = "\n".join(page.get_text() for page in pdf)
    for text in (body, pdf_text):
        assert_clean(text)
        assert text.count("[…]") == 1
        assert text.index("Primer fragmento literal.") < text.index("[…]")
        assert text.index("[…]") < text.index("Segundo fragmento literal.")
        assert "página 91" not in text and "líneas 72" not in text
    assert value == before
    assert value["requests"][0]["review"]["request_snapshot"]["source_refs"] == refs


def test_hidden_citations_still_reject_unknown_references():
    document = build_document(clean_approval(), metadata(), [], "Preparador privado")
    document["requests"][0]["paragraphs"][1]["reference_numbers"] = [999]
    with pytest.raises(DomainError, match="referencias inválidas"):
        _blocks(document)
