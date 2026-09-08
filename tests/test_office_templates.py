"""Membretes ficticios: preservación, firma explícita y validación del pie completo."""

import hashlib
import json
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from PIL import Image
from test_office_renderer import approved, metadata

from pqrs_claro.models import DomainError
from pqrs_claro.office_renderer import build_document, render_office
from pqrs_claro.office_templates import (
    load_letterhead,
    template_profile,
    validate_letterhead,
    verify_preserved_template,
)


def rewrite_member(path, name, change):
    with ZipFile(path) as source:
        items = [(info, source.read(info.filename)) for info in source.infolist()]
    with ZipFile(path, "w") as target:
        for info, value in items:
            target.writestr(info, change(value) if info.filename == name else value)


@pytest.fixture
def template(tmp_path):
    folder = tmp_path / "templates"
    folder.mkdir()
    path = folder / "membrete.docx"
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.3)
    doc.styles["Normal"].font.name = "Arial"
    paragraph = doc.add_paragraph("Atentamente,")
    paragraph.paragraph_format.space_after = Pt(0)
    doc.add_paragraph("NOMBRE ANTIGUO UNO")
    doc.add_paragraph("NOMBRE ANTIGUO DOS")
    section.header.paragraphs[0].text = "Institución ficticia de prueba"
    icon = BytesIO()
    Image.new("RGB", (4, 4), "green").save(icon, format="PNG")
    icon.seek(0)
    section.header.paragraphs[0].add_run().add_picture(icon, width=Inches(0.1))
    section.footer.paragraphs[
        0
    ].text = "Oficina ficticia · Dirección de ejemplo · Teléfono de prueba"
    doc.save(path)
    entry = dict(
        version=1,
        id="ficticio",
        name="Membrete ficticio",
        filename=path.name,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        source_url="https://example.org/template",
    )
    (folder / "active.json").write_text(json.dumps(entry), encoding="utf-8")
    return path, entry, tmp_path


def canonical(template):
    return build_document(
        approved(), metadata(), [], "Preparador", letterhead=load_letterhead(template[2]).public()
    )


def template_pdf(template_path, *, mode="correct"):
    def converter(docx, path):
        profile = template_profile(template_path)
        paragraphs = [p.text for p in Document(docx).paragraphs]
        if mode == "omit_body":
            paragraphs = paragraphs[:-1]
        if mode == "reorder_body":
            paragraphs = list(reversed(paragraphs))
        font = pymupdf.Font("cjk")
        pdf = pymupdf.open()
        page, y = pdf.new_page(width=612, height=792), 130

        def put(page, text, y):
            writer = pymupdf.TextWriter(page.rect)
            writer.append((55, y), text, font=font, fontsize=10)
            writer.write_text(page)

        for text in paragraphs:
            line = ""
            for word in text.split():
                proposed = f"{line} {word}".strip()
                if font.text_length(proposed, fontsize=10) > 490 and line:
                    put(page, line, y)
                    y += 14
                    if y > 690:
                        page, y = pdf.new_page(width=612, height=792), 130
                    line = word
                else:
                    line = proposed
            if line:
                put(page, line, y)
                y += 16
                if y > 690:
                    page, y = pdf.new_page(width=612, height=792), 130
        for page in pdf:
            put(page, profile["regions"][0]["text"], 55)
            if mode != "missing_logo":
                with ZipFile(template_path) as package:
                    image_data = package.read(profile["images"][0]["part"])
                if mode == "changed_logo":
                    altered = BytesIO()
                    Image.new("RGB", (4, 4), "red").save(altered, format="PNG")
                    image_data = altered.getvalue()
                y_image = 300 if mode == "moved_logo" else 70
                page.insert_image(pymupdf.Rect(55, y_image, 62.2, y_image + 7.2), stream=image_data)
                if mode == "extra_logo":
                    page.insert_image(
                        pymupdf.Rect(65, y_image, 72.2, y_image + 7.2), stream=image_data
                    )
            footer = profile["regions"][1]["text"]
            if mode == "changed_footer":
                footer = footer.replace("Oficina", "Entidad diferente")
            if mode != "missing_footer":
                put(page, footer, 755 if mode != "misplaced_footer" else 400)
            if mode == "extra_footer":
                put(page, "Texto adicional que no se puede ocultar en el pie", 780)
        pdf.save(path)
        pdf.close()
        return "Conversor ficticio"

    return converter


def test_registry_absent_means_no_template_and_public_identity_never_contains_path(
    tmp_path, template
):
    assert load_letterhead(tmp_path / "empty") is None
    value = load_letterhead(template[2])
    assert value.path == template[0]
    assert set(value.public()) == {"id", "name", "sha256", "source_url"}
    assert "filename" not in value.public()


@pytest.mark.parametrize(
    "changes",
    [
        {"filename": "../other.docx"},
        {"filename": r"C:\other.docx"},
        {"filename": "other.txt"},
        {"sha256": "a" * 64},
        {"extra": "x"},
        {"source_url": "file:///private"},
        {"version": 2},
    ],
)
def test_registry_fails_closed_for_invalid_configuration_and_changed_original(template, changes):
    path, entry, data = template
    (path.parent / "active.json").write_text(json.dumps({**entry, **changes}), encoding="utf-8")
    with pytest.raises(DomainError):
        load_letterhead(data)


def test_public_identity_is_in_preview_hash_but_cannot_import_local_paths(template):
    first = canonical(template)
    alternative = deepcopy(first["letterhead"])
    alternative["name"] = "Otro nombre de versión"
    changed = build_document(approved(), metadata(), [], "Preparador", letterhead=alternative)
    assert first != changed
    with pytest.raises(DomainError):
        validate_letterhead({**alternative, "path": str(template[0])})


def test_template_preserves_all_protected_parts_geometry_and_uses_only_explicit_signer(template):
    path, _, data = template
    original_sha = hashlib.sha256(path.read_bytes()).hexdigest()
    result = render_office(
        canonical(template), data / "output", template_path=path, converter=template_pdf(path)
    )
    target = data / "output/oficio.docx"
    assert verify_preserved_template(path, target)
    assert result["verification"]["pdf_text_match"]
    assert hashlib.sha256(path.read_bytes()).hexdigest() == original_sha
    doc = Document(target)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "NOMBRE ANTIGUO" not in text
    assert "Firmante ficticio" in text
    assert not text.startswith("Institución ficticia")
    assert doc.styles["Normal"].font.name == "Arial"
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            assert run.font.name == "Times New Roman"
            assert run.font.size.pt == 12
    assert doc.sections[0].top_margin.inches == 1.5
    assert doc.sections[0].left_margin.inches == 0.75


@pytest.mark.parametrize(
    "mode",
    [
        "changed_footer",
        "missing_footer",
        "extra_footer",
        "misplaced_footer",
        "omit_body",
        "reorder_body",
    ],
)
def test_pdf_does_not_hide_missing_changed_or_moved_static_or_approved_text(template, mode):
    path, _, data = template
    with pytest.raises(DomainError):
        render_office(
            canonical(template),
            data / "output",
            template_path=path,
            converter=template_pdf(path, mode=mode),
        )


@pytest.mark.parametrize(
    "member", ["word/header1.xml", "word/footer1.xml", "word/styles.xml", "word/media/image1.png"]
)
def test_protected_ooxml_and_media_changes_are_rejected(template, member):
    path, _, data = template
    render_office(
        canonical(template), data / "output", template_path=path, converter=template_pdf(path)
    )
    target = data / "output/oficio.docx"
    rewrite_member(target, member, lambda value: value + b" ")
    with pytest.raises(DomainError, match="protegida"):
        verify_preserved_template(path, target)


def test_section_change_fails_even_though_document_xml_is_editable(template):
    path, _, data = template
    render_office(
        canonical(template), data / "output", template_path=path, converter=template_pdf(path)
    )
    target = data / "output/oficio.docx"

    def change(value):
        root = etree.fromstring(value)
        root.find(f".//{qn('w:pgMar')}").set(qn("w:top"), "1000")
        return etree.tostring(root)

    rewrite_member(target, "word/document.xml", change)
    with pytest.raises(DomainError, match="geometría"):
        verify_preserved_template(path, target)


def test_template_must_match_preview_and_stay_unchanged_through_conversion(template):
    path, _, data = template
    doc = canonical(template)
    with pytest.raises(DomainError, match="vista previa"):
        render_office(doc, data / "missing", converter=template_pdf(path))
    doc["letterhead"]["sha256"] = "0" * 64
    with pytest.raises(DomainError, match="cambió"):
        render_office(doc, data / "mismatched", template_path=path, converter=template_pdf(path))

    def mutate(docx, pdf):
        template_pdf(path)(docx, pdf)
        path.write_bytes(path.read_bytes() + b"altered")

    with pytest.raises(DomainError, match="cambió"):
        render_office(
            canonical(template), data / "late-change", template_path=path, converter=mutate
        )


def test_choice_fallback_static_text_is_counted_once_and_dynamic_fields_are_rejected(template):
    path = template[0]

    def wrap(value):
        root = etree.fromstring(value)
        alternate = etree.Element(
            "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
        )
        for name in ("Choice", "Fallback"):
            branch = etree.SubElement(
                alternate, "{http://schemas.openxmlformats.org/markup-compatibility/2006}" + name
            )
            branch.append(deepcopy(root[0]))
        root.remove(root[0])
        root.append(alternate)
        return etree.tostring(root)

    rewrite_member(path, "word/footer1.xml", wrap)
    assert template_profile(path)["regions"][1]["text"].count("Oficina ficticia") == 1

    def add_field(value):
        root = etree.fromstring(value)
        etree.SubElement(root, qn("w:instrText")).text = "DATE"
        return etree.tostring(root)

    rewrite_member(path, "word/footer1.xml", add_field)
    with pytest.raises(DomainError, match="estáticos"):
        template_profile(path)


@pytest.mark.parametrize("mode", ["missing_logo", "changed_logo", "moved_logo", "extra_logo"])
def test_pdf_requires_the_same_visible_institutional_image_in_the_header(template, mode):
    path, _, data = template
    with pytest.raises(DomainError, match="imagen|logo"):
        render_office(
            canonical(template),
            data / "output",
            template_path=path,
            converter=template_pdf(path, mode=mode),
        )


def test_only_physical_line_break_after_ascii_hyphen_can_be_rejoined():
    from pqrs_claro.office_renderer import _pdf_text_matches

    expected = "Fecha de conocimiento: 2026-09-07. Referencia siguiente."
    assert _pdf_text_matches(expected, "Fecha de conocimiento: 2026-09-\n07. Referencia siguiente.")
    assert not _pdf_text_matches(
        expected, "Fecha de conocimiento: 2026-09- 07. Referencia siguiente."
    )
    assert not _pdf_text_matches(
        expected, "Fecha de conocimiento: 2026-09\n07. Referencia siguiente."
    )
    assert not _pdf_text_matches(
        expected, "Fecha de conocimiento: 2026-09-\n08. Referencia siguiente."
    )
