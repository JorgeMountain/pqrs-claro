"""Estructura del oficio: texto aprobado intacto y alcance por numeral."""

from copy import deepcopy

import pytest
from docx import Document
from docx.oxml.ns import qn
from pydantic import ValidationError
from test_office_renderer import approved, fake_pdf, metadata, reseal, source

from pqrs_claro.models import DomainError
from pqrs_claro.office_models import OfficeMetadata
from pqrs_claro.office_renderer import (
    _make_template_docx,
    _verify_docx,
    build_document,
    render_office,
)
from pqrs_claro.office_structure import request_fragments
from pqrs_claro.office_templates import template_profile
from pqrs_claro.review_models import ReviewParagraph
from pqrs_claro.review_service import input_hash
from pqrs_claro.source_library import digest


def paragraph(text, section="response", *, chunk_id=None):
    return {
        "text": text,
        "cited_chunk_ids": [chunk_id or source()["chunk_id"]],
        "cited_fact_ids": [],
        "office_section": section,
    }


def structured(*rows):
    value = approved(len(rows))
    for snapshot, paragraphs in zip(value["requests"], rows, strict=True):
        snapshot["review"]["paragraphs"] = deepcopy(paragraphs)
    return reseal(value)


def canonical(value, **changes):
    return build_document(value, metadata(**changes), [], "Preparador de prueba")


def selected(document, kind):
    return [block for block in document["blocks"] if block["kind"] == kind]


def seal_outer_hashes(value):
    """No cambia los hashes de evaluación/aprobación del contenido."""
    for snapshot in value["requests"]:
        snapshot["snapshot_hash"] = digest(
            {key: item for key, item in snapshot.items() if key != "snapshot_hash"}
        )
    value["snapshot_hash"] = digest(
        {key: item for key, item in value.items() if key != "snapshot_hash"}
    )
    return value


def shared_context_document():
    context = paragraph("El antecedente conserva el alcance del período 2015-1.", "context")
    return canonical(
        structured(
            [context, paragraph("La primera respuesta mantiene su alcance.")],
            [context, paragraph("La segunda respuesta mantiene su alcance.")],
        )
    )


def test_sections_follow_letter_order_without_rewriting_approved_paragraphs():
    value = structured(
        [
            paragraph("El soporte se refiere únicamente a 2015-1.", "context"),
            paragraph("No corresponde atribuir ese semestre al segundo docente."),
            paragraph("Esta conclusión se limita al primer numeral.", "conclusion"),
        ],
        [paragraph("El segundo numeral corresponde a 2019-2.")],
    )
    before = deepcopy(value)
    document = canonical(value)
    kinds = [block["kind"] for block in document["blocks"]]
    assert kinds.index("intro") < kinds.index("context") < kinds.index("request")
    responses = [i for i, kind in enumerate(kinds) if kind == "response"]
    assert max(responses) < kinds.index("conclusion") < kinds.index("closing")
    assert len(selected(document, "intro")) == 1
    for snapshot, request in zip(value["requests"], document["requests"], strict=True):
        assert [p["text"] for p in request["paragraphs"]] == [
            p["text"] for p in snapshot["review"]["paragraphs"]
        ]
    assert value == before


def test_shared_context_preserves_every_origin_and_original_position():
    document = shared_context_document()
    contexts = selected(document, "context")
    assert len(contexts) == 1
    assert contexts[0]["request_ids"] == ["P1", "P2"]
    assert contexts[0]["bookmarks"] == ["response_0_0", "response_1_0"]
    assert contexts[0]["bookmark"] == "response_0_0"
    assert [p["bookmark"] for r in document["requests"] for p in r["paragraphs"]] == [
        "response_0_0",
        "response_0_1",
        "response_1_0",
        "response_1_1",
    ]
    assert contexts[0]["reference_numbers"] == [1]
    assert not selected(document, "scope")


@pytest.mark.parametrize("second_text", ["Contexto común. ", "Contexto comun."])
def test_similar_contexts_are_not_deduplicated_by_text_normalization(second_text):
    document = canonical(
        structured(
            [paragraph("Contexto común.", "context"), paragraph("Respuesta uno.")],
            [paragraph(second_text, "context"), paragraph("Respuesta dos.")],
        )
    )
    assert [block["text"] for block in selected(document, "context")] == [
        "Contexto común.",
        second_text,
    ]


def test_identical_context_with_different_passages_keeps_both_citations():
    another = source(
        source_id="source-2",
        document_id="doc-2",
        chunk_id="source-2:1:1-3",
        sha256="d" * 64,
        title="Otro soporte ficticio",
    )
    value = structured(
        [paragraph("El soporte tiene un alcance delimitado.", "context"), paragraph("Uno.")],
        [
            paragraph(
                "El soporte tiene un alcance delimitado.",
                "context",
                chunk_id=another["chunk_id"],
            ),
            paragraph("Dos."),
        ],
    )
    value["requests"][1]["review"]["available_supports"].append(another)
    document = canonical(reseal(value))
    contexts = selected(document, "context")
    assert len(contexts) == 2
    assert [b["request_ids"] for b in contexts] == [["P1"], ["P2"]]
    assert (
        document["requests"][0]["paragraphs"][0]["reference_numbers"]
        != document["requests"][1]["paragraphs"][0]["reference_numbers"]
    )


def test_equal_context_from_request_scoped_facts_is_not_merged():
    value = structured(
        [paragraph("Se conserva la atribución declarada.", "context"), paragraph("Uno.")],
        [paragraph("Se conserva la atribución declarada.", "context"), paragraph("Dos.")],
    )
    for snapshot in value["requests"]:
        row = snapshot["review"]
        fact = {
            "id": "F1",
            "case_id": row["case_id"],
            "request_id": row["request_id"],
            "text": "Declaración ficticia para esta pregunta.",
            "author": "Responsable ficticio",
            "basis": "Consulta ficticia del registro.",
            "observed_at": "2026-09-07",
            "recorded_at": "2026-09-07T22:00:00+00:00",
        }
        fact["sha256"] = digest(fact)
        row["facts"] = [fact]
        row["paragraphs"][0].update(cited_chunk_ids=[], cited_fact_ids=["F1"])
    document = canonical(reseal(value))
    assert len(selected(document, "context")) == 2


def test_equal_responses_and_conclusions_remain_separate_for_each_numeral():
    response = paragraph("Se informa el contenido de este numeral.")
    conclusion = paragraph("Se conserva el límite expuesto.", "conclusion")
    document = canonical(structured([response, conclusion], [response, conclusion]))
    assert [b["text"] for b in selected(document, "response")] == [response["text"]] * 2
    assert [b["text"] for b in selected(document, "conclusion")] == [conclusion["text"]] * 2
    assert len(selected(document, "request")) == 2


def test_identical_context_and_response_remain_in_their_own_sections():
    text = "El antecedente se refiere exclusivamente a 2015-1."
    document = canonical(structured([paragraph(text, "context"), paragraph(text)]))
    assert [b["text"] for b in selected(document, "context")] == [text]
    assert [b["text"] for b in selected(document, "response")] == [text]


@pytest.mark.parametrize("sections", [["context"], ["conclusion"], ["context", "conclusion"]])
def test_context_or_conclusion_never_substitutes_the_answer_to_a_numeral(sections):
    value = structured(
        [paragraph(f"Texto {index}.", section) for index, section in enumerate(sections)]
    )
    with pytest.raises(DomainError):
        canonical(value)


def test_no_optional_conclusion_is_invented_when_none_was_approved():
    document = canonical(structured([paragraph("Respuesta acotada al antecedente.")]))
    assert selected(document, "conclusion") == []
    assert selected(document, "context") == []
    assert len(selected(document, "closing")) == 1


def test_legacy_paragraph_without_office_section_is_a_literal_response():
    value = approved(1)
    assert "office_section" not in value["requests"][0]["review"]["paragraphs"][0]
    document = canonical(value)
    paragraph_value = document["requests"][0]["paragraphs"][0]
    assert paragraph_value["office_section"] == "response"
    assert selected(document, "response")[0]["text"] == paragraph_value["text"]
    assert selected(document, "context") == selected(document, "conclusion") == []


def test_changing_section_after_approval_invalidates_content_hash_even_if_text_is_same():
    value = structured([paragraph("Contexto con cita.", "context"), paragraph("Respuesta.")])
    row = value["requests"][0]["review"]
    old_hash = row["assessment"]["input_hash"]
    row["paragraphs"][0]["office_section"] = "conclusion"
    assert input_hash(row) != old_hash
    with pytest.raises(DomainError):
        canonical(seal_outer_hashes(value))


def test_parent_and_child_keep_order_and_literal_request_quotes():
    value = approved(2)
    identification = deepcopy(value["requests"][0]["identification"])
    literal = [
        "1. Indique cómo se organiza el trámite de inscripción.",
        "1.1. Respecto del numeral anterior, precise el formulario exigido.",
    ]
    for index, request in enumerate(identification["requests"]):
        request["parent_id"] = "P1" if index else None
        request["label"] = "1.1" if index else "1"
        request["source_refs"] = [
            {"page": 1, "line_start": index + 1, "line_end": index + 1, "quote": literal[index]}
        ] * 2
    for snapshot, request in zip(value["requests"], identification["requests"], strict=True):
        snapshot["identification"] = deepcopy(identification)
        row = snapshot["review"]
        row.update({key: request[key] for key in ("label", "title", "description")})
        row["request_snapshot"].update(deepcopy(request))
    document = canonical(reseal(value))
    assert [r["request_id"] for r in document["requests"]] == ["P1", "P2"]
    assert [r["label"] for r in document["requests"]] == ["1", "1.1"]
    assert [b["text"] for b in selected(document, "request")] == literal
    assert len(selected(document, "response")) == 2


def test_received_date_is_optional_explicit_and_invalidates_preview_content():
    value = approved(1)
    without = canonical(value)
    with_date = canonical(value, received_date="2026-09-03")
    assert digest(without) != digest(with_date)
    assert selected(without, "intro") != selected(with_date, "intro")
    with pytest.raises(ValidationError):
        OfficeMetadata.model_validate(metadata(received_date="2026-02-30"))


@pytest.mark.parametrize("section", ["summary", "approval", "", None])
def test_unknown_sections_cannot_be_inserted_as_free_form_export_metadata(section):
    with pytest.raises(ValidationError):
        ReviewParagraph.model_validate(paragraph("Texto citado.", section))


def test_future_nonblocking_action_does_not_prevent_an_otherwise_approved_office():
    value = approved(1)
    value["requests"][0]["review"]["pending"] = [
        {"resolution": None, "stage": "fulfillment", "blocks_approval": False}
    ]
    document = canonical(reseal(value))
    assert len(selected(document, "response")) == 1


@pytest.mark.parametrize(
    "pending",
    [
        {"resolution": None, "stage": "evidence", "blocks_approval": True},
        {"resolution": None, "stage": "fulfillment", "blocks_approval": True},
        {"resolution": None},
    ],
)
def test_unresolved_blocking_checks_still_prevent_export(pending):
    value = approved(1)
    value["requests"][0]["review"]["pending"] = [pending]
    with pytest.raises(DomainError):
        canonical(reseal(value))


def test_shared_context_docx_keeps_all_bookmarks_and_matches_complete_pdf(tmp_path):
    document = shared_context_document()
    result = render_office(document, tmp_path, converter=fake_pdf)
    assert result["verification"]["paragraph_count"] == 4
    assert result["verification"]["docx_exact"] is True
    assert result["verification"]["pdf_text_match"] is True
    _verify_docx(document["blocks"], tmp_path / "oficio.docx")
    word = Document(tmp_path / "oficio.docx")
    context = selected(document, "context")[0]
    matching = [p for p in word.paragraphs if p.text == context["text"]]
    assert len(matching) == 1
    names = [item.get(qn("w:name")) for item in matching[0]._p.findall(qn("w:bookmarkStart"))]
    assert len(names) == len(context["bookmarks"])
    assert set(names) == set(context["bookmarks"])
    starts = word.element.findall(f".//{qn('w:bookmarkStart')}")
    ids = [item.get(qn("w:id")) for item in starts]
    assert len(ids) == len(set(ids)) == 4


def test_template_writer_preserves_multiple_bookmarks_for_shared_context(tmp_path):
    template = tmp_path / "membrete.docx"
    original = Document()
    original.sections[0].header.paragraphs[0].text = "Membrete ficticio"
    original.sections[0].footer.paragraphs[0].text = "Contacto ficticio"
    original.add_paragraph("Atentamente,")
    original.add_paragraph("PERSONA ANTIGUA DE PRUEBA")
    original.save(template)
    document = shared_context_document()
    path = tmp_path / "oficio.docx"
    _make_template_docx(document["blocks"], path, template, template_profile(template))
    assert _verify_docx(document["blocks"], path) is True
    context = selected(document, "context")[0]
    word = Document(path)
    matching = [p for p in word.paragraphs if p.text == context["text"]]
    assert len(matching) == 1
    names = [item.get(qn("w:name")) for item in matching[0]._p.findall(qn("w:bookmarkStart"))]
    assert len(names) == len(context["bookmarks"])
    assert set(names) == set(context["bookmarks"])


@pytest.mark.parametrize("change", ["remove_secondary", "duplicate_secondary"])
def test_verifier_detects_lost_or_duplicated_shared_context_origin(tmp_path, change):
    document = shared_context_document()
    render_office(document, tmp_path, converter=fake_pdf)
    path = tmp_path / "oficio.docx"
    word = Document(path)
    marker = next(
        marker
        for marker in word.element.findall(f".//{qn('w:bookmarkStart')}")
        if marker.get(qn("w:name")) == "response_1_0"
    )
    if change == "remove_secondary":
        marker.getparent().remove(marker)
    else:
        marker.getparent().append(deepcopy(marker))
    word.save(path)
    with pytest.raises(DomainError):
        _verify_docx(document["blocks"], path)


def original_ref(page, start, end, quote):
    return {"page": page, "line_start": start, "line_end": end, "quote": quote}


def test_overlapping_original_quotes_recover_each_original_line_once():
    request = {
        "source_refs": [
            original_ref(1, 1, 3, "Primera frase.\nSegunda frase.\nTercera frase."),
            original_ref(1, 2, 4, "Segunda frase.\nTercera frase.\nCuarta frase."),
        ]
    }
    before = deepcopy(request)
    assert request_fragments(request) == [
        {
            "page": 1,
            "line_start": 1,
            "line_end": 4,
            "text": "Primera frase.\nSegunda frase.\nTercera frase.\nCuarta frase.",
        }
    ]
    assert request == before


def test_discontinuous_quotes_keep_location_and_original_order_independent_of_ref_order():
    refs = [
        original_ref(2, 1, 2, "Página dos.\nContinuación."),
        original_ref(1, 5, 5, "Después del salto."),
        original_ref(1, 1, 2, "Primera línea.\nSegunda línea."),
    ]
    expected = [
        {"page": 1, "line_start": 1, "line_end": 2, "text": "Primera línea.\nSegunda línea."},
        {"page": 1, "line_start": 5, "line_end": 5, "text": "Después del salto."},
        {"page": 2, "line_start": 1, "line_end": 2, "text": "Página dos.\nContinuación."},
    ]
    assert request_fragments({"source_refs": refs}) == expected
    assert request_fragments({"source_refs": list(reversed(refs))}) == expected


def test_equal_words_at_distinct_coordinates_are_never_deduplicated():
    assert request_fragments(
        {
            "source_refs": [
                original_ref(1, 1, 1, "Sí o no."),
                original_ref(1, 2, 2, "Sí o no."),
                original_ref(1, 4, 4, "Sí o no."),
                original_ref(2, 1, 1, "Sí o no."),
            ]
        }
    ) == [
        {"page": 1, "line_start": 1, "line_end": 2, "text": "Sí o no.\nSí o no."},
        {"page": 1, "line_start": 4, "line_end": 4, "text": "Sí o no."},
        {"page": 2, "line_start": 1, "line_end": 1, "text": "Sí o no."},
    ]


def test_contradictory_original_overlap_is_rejected_instead_of_choosing_one_version():
    refs = [
        original_ref(1, 1, 2, "Sí corresponde.\nA este numeral."),
        original_ref(1, 1, 1, "No corresponde."),
    ]
    with pytest.raises(DomainError, match="contradictorios"):
        request_fragments({"source_refs": refs})


def test_fragment_retains_indentation_blank_lines_and_trailing_blank_line():
    text = "  1. Solicitud literal.\n\n    a) Precisión.\n"
    assert request_fragments({"source_refs": [original_ref(3, 8, 11, text)]}) == [
        {"page": 3, "line_start": 8, "line_end": 11, "text": text}
    ]


@pytest.mark.parametrize("identified", [{}, {"source_refs": None}, {"source_refs": []}])
def test_no_original_refs_returns_empty_without_inventing_description_as_quote(identified):
    assert request_fragments(identified) == []


@pytest.mark.parametrize(
    "ref",
    [
        original_ref(True, 1, 1, "Texto."),
        original_ref(1, "1", 1, "Texto."),
        original_ref(1, 1, 1.0, "Texto."),
        original_ref(0, 1, 1, "Texto."),
        original_ref(1, 0, 1, "Texto."),
        original_ref(1, 3, 2, "Texto."),
        original_ref(1, 1, 2, "Una sola línea."),
        original_ref(1, 1, 1, "Dos\nlíneas."),
        original_ref(1, 1, 1, None),
        original_ref(1, 1, 1, " "),
        original_ref(1, 1, 1, "Texto\x00no XML."),
        original_ref(1, 1, 2, "Saltos\r\nalterados."),
        "referencia no estructurada",
    ],
)
def test_unrepresentable_original_reference_is_rejected_without_silent_normalization(ref):
    with pytest.raises(DomainError):
        request_fragments({"source_refs": [ref]})


def test_original_references_must_be_a_list():
    with pytest.raises(DomainError):
        request_fragments({"source_refs": "texto sin coordenadas"})


def approved_with_original_refs(refs):
    value = approved(1)
    snapshot = value["requests"][0]
    snapshot["identification"]["requests"][0]["source_refs"] = deepcopy(refs)
    snapshot["review"]["request_snapshot"]["source_refs"] = deepcopy(refs)
    return reseal(value)


def test_export_of_overlapping_refs_does_not_repeat_original_lines():
    refs = [
        original_ref(1, 1, 3, "Primera frase.\nSegunda frase.\nTercera frase."),
        original_ref(1, 2, 4, "Segunda frase.\nTercera frase.\nCuarta frase."),
    ]
    document = canonical(approved_with_original_refs(refs))
    assert [b["text"] for b in selected(document, "request")] == [
        "Primera frase.\nSegunda frase.\nTercera frase.\nCuarta frase."
    ]
    assert len(selected(document, "response")) == 1


def test_export_discontinuous_refs_keeps_locations_internal_without_creating_more_numerals():
    refs = [
        original_ref(2, 8, 8, "Se precisa la consulta en otro lugar."),
        original_ref(1, 3, 4, "Se formula la consulta.\nSe delimita su objeto."),
    ]
    document = canonical(approved_with_original_refs(refs))
    request_blocks = selected(document, "request")
    assert [b["text"] for b in request_blocks] == [
        "Se formula la consulta.\nSe delimita su objeto.",
        "Se precisa la consulta en otro lugar.",
    ]
    for block, page, start in zip(request_blocks, [1, 2], [3, 8], strict=True):
        assert block["original_location"]["page"] == page
        assert block["original_location"]["line_start"] == start
        assert block["request_ids"] == ["P1"]
    assert [block["text"] for block in selected(document, "omission")] == ["[…]"]
    assert not selected(document, "scope")
    assert len(document["requests"]) == 1
    assert len(selected(document, "response")) == 1


def test_reference_input_order_does_not_change_the_rendered_original_request():
    refs = [original_ref(2, 4, 4, "Segundo fragmento."), original_ref(1, 2, 2, "Primero.")]
    forward = canonical(approved_with_original_refs(refs))
    reversed_refs = canonical(approved_with_original_refs(list(reversed(refs))))
    assert forward["blocks"] == reversed_refs["blocks"]
    assert digest(forward) == digest(reversed_refs)
