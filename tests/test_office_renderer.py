"""Oficios ficticios: misma versión, referencias cerradas y PDF completo en orden."""

import json
import sys
from copy import deepcopy
from types import SimpleNamespace

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from pydantic import ValidationError

from pqrs_claro.models import DomainError
from pqrs_claro.office_models import (
    MAX_ATTACHMENT_BYTES,
    OfficeGenerate,
    OfficeInput,
    OfficeMetadata,
)
from pqrs_claro.office_renderer import build_document, render_office
from pqrs_claro.review_service import input_hash
from pqrs_claro.source_library import digest


def metadata(**changes):
    result = dict(
        institution="Institución ficticia",
        department="Atención ficticia",
        city="Ciudad Ejemplo",
        issue_date="2026-09-07",
        office_number="QA-001",
        petition_number="EJEMPLO-042",
        recipient="Persona ficticia",
        recipient_contact="",
        subject="Respuesta al ejemplo de revisión",
        signer_name="Firmante ficticio",
        signer_role="Cargo ficticio",
    )
    return {**result, **changes}


def source(**changes):
    result = dict(
        source_id="source-1",
        document_id="doc-1",
        version=1,
        chunk_id="source-1:1:1-3",
        title="Regla ficticia",
        filename="regla.txt",
        quote="La inscripción requiere formulario y constancia.",
        page=1,
        line_start=1,
        line_end=3,
        location_kind="section",
        category="institutional_norm",
        role="rule",
        case_id=None,
        sha256="a" * 64,
        validity="current",
    )
    return {**result, **changes}


def approved(count=2):
    requests = [
        dict(
            id=f"P{i + 1}",
            label=str(i + 1),
            title=f"Pregunta {i + 1}",
            description=f"Indique el requisito {i + 1} del ejemplo.",
        )
        for i in range(count)
    ]
    identification = dict(
        subject="Ejercicio ficticio", requester="Persona ficticia", requests=requests
    )
    snapshots = []
    for request in requests:
        review = dict(
            case_id="case-qa",
            request_id=request["id"],
            label=request["label"],
            title=request["title"],
            description=request["description"],
            request_snapshot={**request, "case_id": "case-qa", "allowed_collection_ids": []},
            base_fingerprint="b" * 64,
            stale=False,
            status="ready",
            revision=4,
            paragraphs=[
                dict(
                    text=f"El requisito {request['label']} es presentar el formulario.",
                    cited_chunk_ids=[source()["chunk_id"]],
                    cited_fact_ids=[],
                )
            ],
            available_supports=[
                source(),
                source(
                    source_id="unused",
                    document_id="unused",
                    chunk_id="unused:1",
                    title="Candidato privado no citado",
                ),
            ],
            facts=[],
            pending=[],
            assessment={"status": "passed", "issues": [], "local_checks": []},
        )
        fingerprint = input_hash(review)
        review["assessment"]["input_hash"] = fingerprint
        approval = dict(
            id="approval-" + request["id"],
            approved_by="Persona que revisó, no firmante",
            note="NOTA INTERNA NO PUBLICABLE",
            approved_at="2026-09-07T22:00:00+00:00",
            input_hash=fingerprint,
        )
        snapshot = dict(
            id=approval["id"],
            case_id="case-qa",
            request_id=request["id"],
            case_version=8,
            case_original_sha256="c" * 64,
            identification=deepcopy(identification),
            review=review,
            approval=approval,
        )
        snapshot["snapshot_hash"] = digest(snapshot)
        snapshots.append(snapshot)
    result = dict(
        case_id="case-qa",
        status="approved",
        requests=snapshots,
        approved_count=count,
        total_count=count,
    )
    result["snapshot_hash"] = digest(result)
    return result


def reseal(value):
    for snapshot in value["requests"]:
        fingerprint = input_hash(snapshot["review"])
        snapshot["review"]["assessment"]["input_hash"] = fingerprint
        snapshot["approval"]["input_hash"] = fingerprint
        snapshot["snapshot_hash"] = digest(
            {k: v for k, v in snapshot.items() if k != "snapshot_hash"}
        )
    value["snapshot_hash"] = digest({k: v for k, v in value.items() if k != "snapshot_hash"})
    return value


def document(value=None, attachments=None):
    return build_document(value or approved(), metadata(), attachments or [], "Preparador ficticio")


def fake_pdf(docx_path, pdf_path, *, change=None):
    """Escribe PDF real desde el DOCX real; no simula Microsoft Word ni su maquetación."""
    paragraphs = [p.text for p in Document(docx_path).paragraphs]
    if change:
        paragraphs = change(paragraphs)
    font = pymupdf.Font("cjk")
    pdf = pymupdf.open()
    page = pdf.new_page(width=612, height=792)
    y = 50
    for text in paragraphs:
        line = ""
        for word in text.split():
            proposed = f"{line} {word}" if line else word
            if font.text_length(proposed, fontsize=10) > 510 and line:
                writer = pymupdf.TextWriter(page.rect)
                writer.append((50, y), line, font=font, fontsize=10)
                writer.write_text(page)
                y += 14
                if y > 730:
                    page, y = pdf.new_page(width=612, height=792), 50
                line = word
            else:
                line = proposed
        if line:
            writer = pymupdf.TextWriter(page.rect)
            writer.append((50, y), line, font=font, fontsize=10)
            writer.write_text(page)
            y += 16
            if y > 730:
                page, y = pdf.new_page(width=612, height=792), 50
    pdf.save(pdf_path)
    pdf.close()
    return "Conversor PDF ficticio de pruebas"


def test_canonical_content_first_use_references_and_private_data_exclusion():
    value = approved()
    before = deepcopy(value)
    doc = document(value)
    assert value == before
    assert list(doc) == ["metadata", "requests", "references", "attachments", "prepared_by"]
    assert len(doc["references"]) == 1
    assert [r["request_id"] for r in doc["requests"]] == ["P1", "P2"]
    assert all(r["paragraphs"][0]["reference_numbers"] == [1] for r in doc["requests"])
    assert doc["metadata"]["signer_name"] == "Firmante ficticio"
    assert "NOTA INTERNA" not in str(doc)
    assert "Persona que revisó" not in str(doc)
    assert "Candidato privado" not in str(doc)
    assert "base_fingerprint" not in str(doc)


def test_fact_reference_preserves_attribution_without_publishing_internal_basis():
    value = approved(1)
    row = value["requests"][0]["review"]
    fact = dict(
        id="F1",
        case_id="case-qa",
        request_id="P1",
        text="El trámite sigue en revisión.",
        author="Responsable ficticio",
        basis="NOTA INTERNA sobre la consulta realizada.",
        observed_at="2026-09-07",
        recorded_at="2026-09-07T22:00:00+00:00",
    )
    fact["sha256"] = digest(fact)
    row["facts"] = [fact]
    row["paragraphs"][0].update(cited_chunk_ids=[], cited_fact_ids=["F1"])
    doc = document(reseal(value))
    assert doc["references"] == [
        {
            "number": 1,
            "kind": "fact",
            "text": "Declaración de Responsable ficticio. Fecha de conocimiento: 2026-09-07.",
        }
    ]
    assert "NOTA INTERNA" not in str(doc)


@pytest.mark.parametrize(
    "mode",
    [
        "global_hash",
        "snapshot_hash",
        "review_hash",
        "missing",
        "order",
        "case",
        "pending",
        "title",
        "fact_scope",
    ],
)
def test_incomplete_changed_or_cross_case_approval_is_rejected(mode):
    value = approved()
    row = value["requests"][0]["review"]
    if mode == "global_hash":
        value["snapshot_hash"] = "0" * 64
    elif mode == "snapshot_hash":
        value["requests"][0]["snapshot_hash"] = "0" * 64
        value["snapshot_hash"] = digest({k: v for k, v in value.items() if k != "snapshot_hash"})
    elif mode == "review_hash":
        row["paragraphs"][0]["text"] = "Cambió después de aprobar."
        for snapshot in value["requests"]:
            snapshot["snapshot_hash"] = digest(
                {k: v for k, v in snapshot.items() if k != "snapshot_hash"}
            )
        value["snapshot_hash"] = digest({k: v for k, v in value.items() if k != "snapshot_hash"})
    else:
        if mode == "missing":
            value["requests"].pop()
        elif mode == "order":
            value["requests"].reverse()
        elif mode == "case":
            row["case_id"] = "another-case"
        elif mode == "pending":
            row["pending"] = [{"resolution": None}]
        elif mode == "title":
            row["title"] = "Otro asunto"
        elif mode == "fact_scope":
            row["paragraphs"][0]["cited_fact_ids"] = ["foreign-fact"]
        reseal(value)
    with pytest.raises(DomainError):
        document(value)


def test_unknown_source_and_mutated_fact_are_rejected_even_with_valid_outer_hashes():
    value = approved(1)
    value["requests"][0]["review"]["paragraphs"][0]["cited_chunk_ids"] = ["unknown"]
    with pytest.raises(DomainError, match="desconocida"):
        document(reseal(value))
    value = approved(1)
    row = value["requests"][0]["review"]
    row["facts"] = [dict(id="F1", case_id="case-qa", request_id="P1", sha256="wrong")]
    row["paragraphs"][0]["cited_fact_ids"] = ["F1"]
    with pytest.raises(DomainError, match="integridad"):
        document(reseal(value))


@pytest.mark.parametrize(
    "field,value",
    [
        ("institution", "  "),
        ("signer_name", ""),
        ("subject", "x\x00y"),
        ("recipient", "x\ud800y"),
        ("issue_date", "2026-02-30"),
    ],
)
def test_metadata_rejects_missing_names_invalid_dates_and_xml_controls(field, value):
    data = metadata()
    data[field] = value
    with pytest.raises(ValidationError):
        OfficeMetadata.model_validate(data)


def test_input_extra_fields_duplicate_attachments_and_preview_hash_are_closed():
    data = dict(approval_snapshot_hash="a" * 64, metadata=metadata(), prepared_by="QA")
    with pytest.raises(ValidationError):
        OfficeInput.model_validate({**data, "approval": True})
    with pytest.raises(ValidationError):
        OfficeInput.model_validate({**data, "attachment_source_ids": ["id", "id"]})
    with pytest.raises(ValidationError):
        OfficeGenerate.model_validate({**data, "preview_hash": "invalid"})


def test_attachments_only_cited_exact_sources_and_aggregate_limit():
    item = dict(
        source_id="source-1",
        title="Regla ficticia",
        filename="regla.txt",
        size=100,
        sha256="a" * 64,
    )
    assert document(attachments=[item])["attachments"] == [item]
    for changes in (
        {"source_id": "unused"},
        {"sha256": "b" * 64},
        {"filename": "../regla.txt"},
        {"size": MAX_ATTACHMENT_BYTES + 1},
    ):
        with pytest.raises(DomainError):
            document(attachments=[{**item, **changes}])
    with pytest.raises(DomainError):
        document(attachments=[item, item])


def test_render_preserves_word_paragraphs_and_builds_real_pdf_with_same_text(tmp_path):
    canonical = document()
    result = render_office(canonical, tmp_path, converter=fake_pdf)
    assert result["verification"] == dict(
        docx_exact=True, pdf_text_match=True, request_count=2, paragraph_count=2
    )
    texts = [p.text for p in Document(tmp_path / "oficio.docx").paragraphs]
    assert canonical["requests"][0]["paragraphs"][0]["text"] in texts
    assert "Referencias: [1]." in texts
    assert "Firmante ficticio" in texts
    assert all("Persona que revisó" not in p for p in texts)
    assert result["page_count"] >= 1


def test_word_styles_remove_inherited_theme_fonts_and_blue_title_border(tmp_path):
    render_office(document(), tmp_path, converter=fake_pdf)
    docx = Document(tmp_path / "oficio.docx")
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Footer"):
        style = docx.styles[name]
        fonts = style.element.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
        assert all("theme" not in attribute.casefold() for attribute in fonts.attrib)
        assert all(
            fonts.get(qn(f"w:{attribute}")) == "Times New Roman"
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs")
        )
        assert style.element.find(f"{qn('w:pPr')}/{qn('w:pBdr')}") is None
        assert style.element.find(f"{qn('w:rPr')}/{qn('w:spacing')}") is None
        assert str(style.font.color.rgb) == "000000"
    assert docx.styles["Title"].font.size.pt == 15
    assert docx.styles["Heading 1"].font.size.pt == 12


@pytest.mark.parametrize(
    "mode", ["correct", "wrong_number", "extra_bottom", "body_number", "uncentered"]
)
def test_pdf_excludes_only_exact_centered_page_number_and_keeps_other_text(tmp_path, mode):
    def converter(docx, path):
        fake_pdf(docx, path)
        font = pymupdf.Font("tiro")
        with pymupdf.open(path) as pdf:
            for page in pdf:
                text = f"Página {page.number + 1 + (mode == 'wrong_number')}"
                x = (page.rect.width - font.text_length(text, fontsize=12)) / 2
                page.insert_text(
                    (x if mode != "uncentered" else 50, 758), text, fontsize=12, fontname="tiro"
                )
                if mode == "extra_bottom":
                    page.insert_text((50, 780), "Un dato adicional no aprobado.", fontname="tiro")
                if mode == "body_number":
                    page.insert_text((x, 720), "Página 1", fontname="tiro")
            temporary = path.with_suffix(".converted.pdf")
            pdf.save(temporary)
        temporary.replace(path)

    if mode == "correct":
        result = render_office(document(), tmp_path, converter=converter)
        assert result["verification"]["pdf_text_match"] is True
    else:
        with pytest.raises(DomainError, match="bloques y referencias"):
            render_office(document(), tmp_path, converter=converter)


def test_unicode_line_breaks_and_multi_page_response_are_preserved(tmp_path):
    value = approved(1)
    row = value["requests"][0]["review"]
    text = "Español áéíóú ñ; № 42, α — intervalo 2010–2015.\nUna segunda línea\ty un tabulador. "
    row["paragraphs"] = [
        {"text": text * 25, "cited_chunk_ids": [source()["chunk_id"]], "cited_fact_ids": []}
        for _ in range(4)
    ]
    canonical = document(reseal(value))
    result = render_office(canonical, tmp_path, converter=fake_pdf)
    assert result["page_count"] > 1
    texts = [p.text for p in Document(tmp_path / "oficio.docx").paragraphs]
    assert texts.count(text * 25) == 4


@pytest.mark.parametrize("mode", ["omit", "reorder", "replace", "append", "empty", "invalid"])
def test_pdf_integrity_rejects_missing_changed_reordered_and_extra_text(tmp_path, mode):
    def converter(docx, pdf):
        if mode == "empty":
            return None
        if mode == "invalid":
            pdf.write_bytes(b"not a PDF")
            return None

        def change(paragraphs):
            if mode == "omit":
                return paragraphs[:-1]
            if mode == "reorder":
                return list(reversed(paragraphs))
            if mode == "append":
                return [*paragraphs, "Se aprobó un hecho inventado."]
            return [p.replace("formulario", "pago") for p in paragraphs]

        return fake_pdf(docx, pdf, change=change)

    with pytest.raises(DomainError):
        render_office(document(), tmp_path, converter=converter)


def test_converter_must_not_modify_the_canonical_docx(tmp_path):
    def converter(docx, pdf):
        changed = Document(docx)
        changed.paragraphs[0].text = "Contenido alterado"
        changed.save(docx)
        return fake_pdf(docx, pdf)

    with pytest.raises(DomainError, match="modificó el Word"):
        render_office(document(), tmp_path, converter=converter)


def test_renderer_rejects_old_files_and_provider_errors_without_raw_details(tmp_path):
    def converter(*args):
        raise RuntimeError("PRIVATE PATH SECRET")

    with pytest.raises(DomainError, match="no quedó listo") as failure:
        render_office(document(), tmp_path, converter=converter)
    assert "PRIVATE" not in str(failure.value)
    with pytest.raises(DomainError, match="sobrescribe"):
        render_office(document(), tmp_path, converter=fake_pdf)


def test_document_and_pdf_limits_never_truncate_silently(tmp_path, monkeypatch):
    from pqrs_claro import office_renderer

    monkeypatch.setattr(office_renderer, "MAX_OFFICE_CHARS", 10)
    with pytest.raises(DomainError, match="600.000"):
        document()
    monkeypatch.setattr(office_renderer, "MAX_OFFICE_CHARS", 600000)

    def too_many_pages(docx, path):
        pdf = pymupdf.open()
        for _ in range(201):
            pdf.new_page()
        pdf.save(path)
        pdf.close()

    with pytest.raises(DomainError, match="200 páginas"):
        render_office(document(), tmp_path, converter=too_many_pages)


def test_capability_reports_missing_or_unsupported_converter(monkeypatch):
    from pqrs_claro import office_word

    monkeypatch.setenv("PQRS_CLARO_PDF_ENGINE", "unsupported")
    assert not office_word.pdf_capabilities()["pdf_available"]
    monkeypatch.setenv("PQRS_CLARO_PDF_ENGINE", "word")
    monkeypatch.setattr(office_word.sys, "platform", "linux")
    assert "Windows" in office_word.pdf_capabilities()["error"]
    monkeypatch.setattr(office_word.sys, "platform", "win32")
    monkeypatch.setattr(office_word, "_word_installed", lambda: False)
    assert "Microsoft Word" in office_word.pdf_capabilities()["error"]


def test_word_timeout_terminates_only_the_owned_worker_target(tmp_path, monkeypatch):
    from pqrs_claro import office_word

    docx = tmp_path / "canonical.docx"
    docx.write_bytes(b"fake fixture")
    monkeypatch.setattr(office_word, "pdf_capabilities", lambda: {"pdf_available": True})
    killed, targets, commands = [], [], []

    class Process:
        def __init__(self):
            self.calls = 0

        def communicate(self, timeout=None):
            self.calls += 1
            if self.calls == 1:
                raise office_word.subprocess.TimeoutExpired("test", timeout)
            return b"", b""

        def kill(self):
            killed.append(True)

    def popen(command, **kwargs):
        commands.append(command)
        return Process()

    monkeypatch.setattr(office_word.subprocess, "Popen", popen)
    monkeypatch.setattr(office_word, "_terminate_owned_word", targets.append)
    with pytest.raises(DomainError, match="tiempo máximo"):
        office_word.convert_word(docx, tmp_path / "oficio.pdf")
    assert killed == [True]
    assert targets == [tmp_path / ".word-worker.json"]
    assert commands[0][1:3] == ["-m", "pqrs_claro.office_word"]


def test_word_missing_output_is_not_reported_as_pdf_ready(tmp_path, monkeypatch):
    from pqrs_claro import office_word

    docx = tmp_path / "canonical.docx"
    docx.write_bytes(b"fake fixture")
    monkeypatch.setattr(office_word, "pdf_capabilities", lambda: {"pdf_available": True})
    monkeypatch.setattr(office_word, "_terminate_owned_word", lambda _: None)
    monkeypatch.setattr(
        office_word.subprocess,
        "Popen",
        lambda *args, **kwargs: SimpleNamespace(
            communicate=lambda **kwargs: (b"", b""), returncode=0
        ),
    )
    with pytest.raises(DomainError, match="no pudo convertir"):
        office_word.convert_word(docx, tmp_path / "oficio.pdf")


@pytest.mark.parametrize(
    "owned,created,should_kill",
    [(False, "original", False), (True, "new", False), (True, "original", True)],
)
def test_cleanup_checks_ownership_and_process_creation_time(
    tmp_path, monkeypatch, owned, created, should_kill
):
    from pqrs_claro.office_word import _terminate_owned_word

    control = tmp_path / "worker.json"
    control.write_text(json.dumps(dict(owned=owned, pid=42, process_created="original")))
    opened, terminated, closed = [], [], []
    handle = SimpleNamespace(Close=lambda: closed.append(True))

    def open_process(flags, inherit, pid):
        opened.append(pid)
        return handle

    monkeypatch.setitem(
        sys.modules,
        "win32api",
        SimpleNamespace(
            OpenProcess=open_process, TerminateProcess=lambda h, code: terminated.append(h)
        ),
    )
    monkeypatch.setitem(
        sys.modules, "win32con", SimpleNamespace(PROCESS_QUERY_INFORMATION=1, PROCESS_TERMINATE=2)
    )
    monkeypatch.setitem(
        sys.modules,
        "win32process",
        SimpleNamespace(GetProcessTimes=lambda h: dict(CreationTime=created)),
    )
    _terminate_owned_word(control)
    assert bool(terminated) is should_kill
    assert opened == ([42] if owned else [])
    assert closed == ([True] if owned else [])


def test_cleanup_preserves_original_failure_if_word_already_exited(tmp_path, monkeypatch):
    from pqrs_claro.office_word import _terminate_owned_word

    class ComProcessError(Exception):
        pass

    def disappeared(*args):
        raise ComProcessError("The process has already exited")

    control = tmp_path / "worker.json"
    control.write_text(json.dumps(dict(owned=True, pid=42, process_created="original")))
    monkeypatch.setitem(sys.modules, "win32api", SimpleNamespace(OpenProcess=disappeared))
    monkeypatch.setitem(
        sys.modules, "win32con", SimpleNamespace(PROCESS_QUERY_INFORMATION=1, PROCESS_TERMINATE=2)
    )
    monkeypatch.setitem(sys.modules, "win32process", SimpleNamespace())
    _terminate_owned_word(control)
